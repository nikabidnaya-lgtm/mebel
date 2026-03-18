import asyncio
import hashlib
import json
import re
from collections import Counter, defaultdict
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from urllib.parse import urljoin

import httpx
import imagehash
import pandas as pd
from dotenv import load_dotenv
from loguru import logger
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image
from playwright.async_api import BrowserContext, Page, async_playwright
from rapidfuzz import fuzz

# ================= НАСТРОЙКИ =================
EXCEL_PATH = r"C:\Users\BAand\Downloads\Перечень мебели (3).xlsx"
OUTPUT_EXCEL = "Аналоги_мебель_с_фото_3сайта.xlsx"
DEBUG_ANALOGS_DUMP = "debug_all_analogs.xlsx"
CACHE_DIR = Path("cache")
CATEGORY_CACHE_DIR = CACHE_DIR / "categories"
IMAGE_CACHE_DIR = CACHE_DIR / "images"
CATEGORY_PAGE_LIMIT = 2
MAX_LINKS_PER_PAGE = 36
DEFAULT_ANALOGS_PER_ITEM = 3
TEXT_CANDIDATES_PER_ITEM = 20
RELEVANCE_POOL_FOR_CHEAPEST = 12
CATEGORY_CONCURRENCY = 3
PRODUCT_CONCURRENCY = 6
IMAGE_SCORE_WEIGHT = 0.35
TEXT_SCORE_WEIGHT = 0.65
REQUEST_TIMEOUT_MS = 45000
EMBED_IMAGE_WIDTH = 120
EMBED_IMAGE_HEIGHT = 120
load_dotenv()

logger.add("scraper.log", rotation="10 MB", level="INFO", encoding="utf-8")
logger.info("🔧 Логирование запущено — теперь всё будет видно в реальном времени!")

CATEGORY_RULES: Dict[str, Dict[str, Any]] = {
    "chair": {
        "label": "Кресла и стулья",
        "keywords": ["кресл", "стул", "chair", "кресло", "стуль", "табурет", "пуф"],
        "sites": {
            "profoffice": ["/catalog/chairs/"],
            "riva": ["/catalog/kresla/"],
            "attento": ["/kresla/", "/stulya-ofisnye/"]
        },
    },
    "soft": {
        "label": "Мягкая мебель",
        "keywords": ["диван", "соф", "банкет", "пуф", "мягк", "soft"],
        "sites": {
            "profoffice": ["/catalog/waiting_areas/"],
            "riva": ["/catalog/myagkaya-mebel/"],
            "attento": ["/myagkaya-mebel/"]
        },
    },
    "reception": {
        "label": "Ресепшн и переговорные",
        "keywords": ["ресеп", "стойк", "reception", "переговор", "conference"],
        "sites": {
            "profoffice": ["/catalog/offices/"],
            "riva": ["/catalog/stoyki-resepshn/", "/catalog/mebel-dlya-peregovornykh-zon/"],
            "attento": ["/stoly-dlya-peregovorov/"]
        },
    },
    "table": {
        "label": "Столы",
        "keywords": ["стол", "table", "тумб", "приставк", "журнальн", "кофейный"],
        "sites": {
            "profoffice": ["/catalog/offices/"],
            "riva": ["/catalog/mebel-dlya-personala/", "/catalog/mebel-dlya-peregovornykh-zon/"],
            "attento": ["/ofisnye-stoly/", "/stoly-dlya-peregovorov/"]
        },
    },
    "storage": {
        "label": "Шкафы и хранение",
        "keywords": ["шкаф", "стеллаж", "тумб", "комод", "ящик", "гардероб", "кухон"],
        "sites": {
            "profoffice": ["/catalog/offices/"],
            "riva": ["/catalog/mebel-dlya-personala/"],
            "attento": ["/operativnaya-mebel/", "/kabinety-rukovoditeley/"]
        },
    },
    "office": {
        "label": "Офисная мебель",
        "keywords": ["офис", "кабинет", "мебел", "system", "система", "перегород", "кашпо"],
        "sites": {
            "profoffice": ["/catalog/offices/"],
            "riva": ["/catalog/mebel-dlya-personala/"],
            "attento": ["/operativnaya-mebel/", "/kabinety-rukovoditeley/"]
        },
    },
}

SITE_CONFIGS: Dict[str, Dict[str, Any]] = {
    "profoffice": {
        "base": "https://www.profoffice.ru",
        "product_link_selectors": [".catalog-item a[href]", ".item a[href]", "a[href*='/catalog/'][href$='/']"],
        "product_url_patterns": [re.compile(r"/catalog/.+/.+/.+/"), re.compile(r"/catalog/.+/.+_/")],
        "deny_url_patterns": [re.compile(r"/catalog/[^/]+/$"), re.compile(r"/catalog/.+/filter/")],
        "title_selectors": ["h1", ".bx-title", ".detail-title"],
        "description_selectors": [".detail-text", ".bx-detail-description", ".content", "main"],
        "price_selectors": [
            "[itemprop='price']",
            "meta[itemprop='price']",
            "meta[property='product:price:amount']",
            ".price",
            ".price-item",
            ".price-current",
            "[class*='price']",
            "[data-price]",
        ],
        "image_selectors": [".detail img", ".product-card img", "img[itemprop='image']", "img"],
    },
    "riva": {
        "base": "https://riva.ru",
        "product_link_selectors": [".catalog-section a[href]", ".catalog-list a[href]", "a[href*='/catalog/'][href$='/']"],
        "product_url_patterns": [re.compile(r"/catalog/.+/.+/")],
        "deny_url_patterns": [re.compile(r"/catalog/[^/]+/$"), re.compile(r"/catalog/.+/filter/")],
        "title_selectors": ["h1", ".product-name", ".catalog-element__title"],
        "description_selectors": [".catalog-element__text", ".text-block", "[itemprop='description']", "main"],
        "price_selectors": [
            "[itemprop='price']",
            "meta[itemprop='price']",
            "meta[property='product:price:amount']",
            ".price",
            ".catalog-element__price",
            ".price-current",
            "[class*='price']",
            "[data-price]",
        ],
        "image_selectors": [".catalog-element img", "img[itemprop='image']", "img"],
    },
    "attento": {
        "base": "https://attento.ru",
        "product_link_selectors": [".products a[href*='/product/']", ".catalog a[href*='/product/']", "a[href*='/product/']"],
        "product_url_patterns": [re.compile(r"/product/[^/?#]+/?")],
        "deny_url_patterns": [],
        "title_selectors": ["h1", ".product_title", ".entry-title"],
        "description_selectors": [".woocommerce-product-details__short-description", ".product-description", "main"],
        "price_selectors": [
            ".price",
            ".woocommerce-Price-amount",
            "meta[property='product:price:amount']",
            "meta[itemprop='price']",
            "[class*='price']",
            "[data-price]",
        ],
        "image_selectors": [".woocommerce-product-gallery img", "img.wp-post-image", "img"],
    },
}

STOP_WORDS = {"тип", "для", "и", "с", "со", "на", "по", "из", "под", "над", "без", "мебель", "комплект"}

try:
    import psutil

    def get_optimal_workers() -> Tuple[int, int]:
        cpu = psutil.cpu_count(logical=True)
        usage = psutil.cpu_percent(interval=0.4)
        ram_free_gb = psutil.virtual_memory().available / (1024**3)
        category_workers = max(2, min(4, int(cpu * 0.25)))
        product_workers = max(4, min(10, int(cpu * 0.55)))
        if usage > 60:
            category_workers = max(2, category_workers - 1)
            product_workers = max(4, product_workers - 2)
        if ram_free_gb < 4:
            category_workers, product_workers = 2, 4
        logger.info(
            f"💻 Авто-определение потоков: категории={category_workers}, карточки={product_workers} "
            f"(CPU {cpu}, загрузка {usage:.1f}%, RAM свободно {ram_free_gb:.1f} ГБ)"
        )
        return category_workers, product_workers
except ImportError:
    def get_optimal_workers() -> Tuple[int, int]:
        logger.warning("⚠️ psutil не установлен → используем 3 потока категорий и 6 потоков карточек")
        return 3, 6


def ask_analog_count() -> int:
    raw = input("Сколько аналогов подобрать для каждого товара? Введите 1, 2 или 3 [по умолчанию 3]: ").strip()
    if raw in {"1", "2", "3"}:
        value = int(raw)
        logger.info(f"🧾 Пользователь выбрал число аналогов: {value}")
        return value
    logger.warning(f"⚠️ Некорректный ввод '{raw or 'пусто'}' — используем значение по умолчанию: {DEFAULT_ANALOGS_PER_ITEM}")
    return DEFAULT_ANALOGS_PER_ITEM


# ================= ТЕКСТ, КАТЕГОРИИ, КЭШ =================
def normalize_text(text: str) -> str:
    text = re.sub(r"[^a-zA-Zа-яА-Я0-9]+", " ", (text or "").lower())
    return re.sub(r"\s+", " ", text).strip()


def tokenize(text: str) -> List[str]:
    return [token for token in normalize_text(text).split() if len(token) > 2 and token not in STOP_WORDS]


def classify_item_category(name: str, description: str) -> str:
    source = normalize_text(f"{name} {description}")
    scores: Dict[str, int] = {}
    for category_key, rule in CATEGORY_RULES.items():
        score = sum(3 for keyword in rule["keywords"] if keyword in source)
        score += fuzz.partial_ratio(category_key, source) // 25
        scores[category_key] = score
    best_category = max(scores, key=scores.get)
    if scores[best_category] == 0:
        return "office"
    return best_category


def build_scrape_plan(df: pd.DataFrame) -> Tuple[Dict[str, List[str]], Dict[Tuple[str, str], List[str]]]:
    needed_categories = sorted(set(df["predicted_category"].tolist()))
    plan: Dict[str, List[str]] = defaultdict(list)
    path_to_categories: Dict[Tuple[str, str], List[str]] = defaultdict(list)
    for category_key in needed_categories:
        for site_key, paths in CATEGORY_RULES[category_key]["sites"].items():
            for path in paths:
                plan[site_key].append(path)
                path_to_categories[(site_key, path)].append(category_key)
    for site_key, paths in plan.items():
        plan[site_key] = sorted(set(paths))
    for key, categories in path_to_categories.items():
        path_to_categories[key] = sorted(set(categories))
    return dict(plan), dict(path_to_categories)


def ensure_cache_dirs() -> None:
    CATEGORY_CACHE_DIR.mkdir(parents=True, exist_ok=True)
    IMAGE_CACHE_DIR.mkdir(parents=True, exist_ok=True)


def cache_key(*parts: str) -> str:
    return hashlib.md5("||".join(parts).encode("utf-8")).hexdigest()


def category_cache_path(site_key: str, cat_path: str) -> Path:
    return CATEGORY_CACHE_DIR / f"{site_key}_{cache_key(cat_path)}.json"


def load_category_cache(site_key: str, cat_path: str) -> Optional[List[Dict[str, Any]]]:
    cache_file = category_cache_path(site_key, cat_path)
    if not cache_file.exists():
        return None
    try:
        data = json.loads(cache_file.read_text(encoding="utf-8"))
        logger.info(f"♻️ Загружен кэш категории {site_key}{cat_path}: {len(data)} товаров")
        return data
    except Exception as exc:
        logger.warning(f"⚠️ Не удалось прочитать кэш {cache_file}: {exc}")
        return None


def save_category_cache(site_key: str, cat_path: str, products: List[Dict[str, Any]]) -> None:
    category_cache_path(site_key, cat_path).write_text(json.dumps(products, ensure_ascii=False, indent=2), encoding="utf-8")


# ================= ФОТО-КЭШ =================
IMAGE_HASH_CACHE: Dict[str, imagehash.ImageHash] = {}
IMAGE_BYTES_CACHE: Dict[str, Path] = {}


def local_image_cache_path(url: str) -> Path:
    suffix = Path(url.split("?")[0]).suffix or ".img"
    return IMAGE_CACHE_DIR / f"{cache_key(url)}{suffix}"


def download_to_cache(url: str) -> Optional[Path]:
    if not url:
        return None
    if url in IMAGE_BYTES_CACHE:
        return IMAGE_BYTES_CACHE[url]
    file_path = local_image_cache_path(url)
    if file_path.exists():
        IMAGE_BYTES_CACHE[url] = file_path
        return file_path
    try:
        with httpx.Client(timeout=15, follow_redirects=True, headers={"User-Agent": "Mozilla/5.0"}) as client:
            response = client.get(url)
            if response.status_code != 200:
                return None
            file_path.write_bytes(response.content)
            IMAGE_BYTES_CACHE[url] = file_path
            return file_path
    except Exception:
        return None


def get_image_hash(path_or_url: str, is_url: bool = False) -> Optional[imagehash.ImageHash]:
    if not path_or_url:
        return None
    cache_id = f"url:{path_or_url}" if is_url else f"file:{path_or_url}"
    if cache_id in IMAGE_HASH_CACHE:
        return IMAGE_HASH_CACHE[cache_id]
    try:
        image_path = download_to_cache(path_or_url) if is_url else Path(path_or_url)
        if not image_path or not Path(image_path).exists():
            return None
        img_hash = imagehash.phash(Image.open(image_path))
        IMAGE_HASH_CACHE[cache_id] = img_hash
        return img_hash
    except Exception:
        return None


def image_similarity_score(original_path: str, analog_url: str) -> float:
    original_hash = get_image_hash(original_path, is_url=False)
    analog_hash = get_image_hash(analog_url, is_url=True)
    if not original_hash or not analog_hash:
        return 0.0
    return max(0.0, 1 - (original_hash - analog_hash) / 64.0)


# ================= ПАРСИНГ =================
def clean_price(text: str) -> int:
    digits = re.sub(r"[^\d]", "", str(text or ""))
    return int(digits) if digits else 0


def extract_price_candidates(text: str) -> List[int]:
    if not text:
        return []
    patterns = [
        r'content=["\'](\d{3,9})["\']',
        r'price["\'\s:>=]+(\d{3,9})',
        r'(\d{1,3}(?:[\s\u00A0]?\d{3}){1,3})\s*(?:₽|руб\.?|RUB)',
        r'(?:₽|руб\.?|RUB)\s*(\d{1,3}(?:[\s\u00A0]?\d{3}){1,3})',
        r'data-price=["\'](\d{3,9})["\']',
    ]
    values: List[int] = []
    for pattern in patterns:
        for match in re.findall(pattern, text, flags=re.I):
            value = clean_price(match)
            if value >= 100:
                values.append(value)
    return values


async def first_text(page: Page, selectors: List[str], timeout: int = 5000) -> str:
    for selector in selectors:
        try:
            locator = page.locator(selector)
            if await locator.count() > 0:
                text = (await locator.first.inner_text(timeout=timeout)).strip()
                if text:
                    return text
        except Exception:
            continue
    return ""


async def first_attr(page: Page, selectors: List[str], attr: str, timeout: int = 5000) -> str:
    for selector in selectors:
        try:
            locator = page.locator(selector)
            if await locator.count() > 0:
                value = await locator.first.get_attribute(attr, timeout=timeout)
                if value:
                    return value.strip()
        except Exception:
            continue
    return ""


async def create_context(pw):
    logger.info("🌐 Запуск браузера Playwright (headless)...")
    browser = await pw.chromium.launch(headless=True)
    context = await browser.new_context(
        viewport={"width": 1600, "height": 1200},
        locale="ru-RU",
        user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/134.0.0.0 Safari/537.36",
    )
    await context.add_init_script("""() => { Object.defineProperty(navigator, 'webdriver', { get: () => undefined }); }""")
    logger.info("✅ Браузерный контекст создан")
    return context


async def collect_product_urls(page: Page, site_key: str, cat_path: str) -> List[str]:
    config = SITE_CONFIGS[site_key]
    seen = set()
    found: List[str] = []
    for selector in config["product_link_selectors"]:
        try:
            hrefs = await page.locator(selector).evaluate_all("els => els.map(el => el.href || el.getAttribute('href') || '')")
        except Exception:
            continue
        for href in hrefs:
            if not href:
                continue
            full_url = urljoin(config["base"], href)
            lower = full_url.lower()
            if full_url in seen:
                continue
            if any(token in lower for token in ["filter", "sort", "login", "cart", "search", "javascript:", "#"]):
                continue
            if any(pattern.search(lower) for pattern in config["deny_url_patterns"]):
                continue
            if not any(pattern.search(lower) for pattern in config["product_url_patterns"]):
                continue
            if lower.rstrip("/") == urljoin(config["base"], cat_path).lower().rstrip("/"):
                continue
            seen.add(full_url)
            found.append(full_url)
    return found[:MAX_LINKS_PER_PAGE]


async def extract_price(page: Page, config: Dict[str, Any]) -> Tuple[int, str]:
    tried: List[str] = []
    for selector in config["price_selectors"]:
        try:
            locator = page.locator(selector)
            if await locator.count() == 0:
                continue
            value = await locator.first.get_attribute("content")
            if value and clean_price(value) > 0:
                return clean_price(value), f"attr:content:{selector}"
            value = await locator.first.get_attribute("data-price")
            if value and clean_price(value) > 0:
                return clean_price(value), f"attr:data-price:{selector}"
            text = (await locator.first.inner_text(timeout=3000)).strip()
            if clean_price(text) > 0:
                return clean_price(text), f"text:{selector}"
            tried.append(selector)
        except Exception:
            tried.append(selector)
            continue

    html = await page.content()
    body_text = await first_text(page, ["body"], timeout=3000)
    for source_name, source in [("html", html), ("body", body_text)]:
        candidates = extract_price_candidates(source)
        if candidates:
            return min(candidates), f"regex:{source_name}"

    meta_price = await first_attr(page, ["meta[property='product:price:amount']", "meta[itemprop='price']"], "content")
    if clean_price(meta_price) > 0:
        return clean_price(meta_price), "meta-content"

    return 0, f"unparsed:{tried}"


async def scrape_product_page(page: Page, site_key: str, url: str) -> Optional[Dict[str, Any]]:
    config = SITE_CONFIGS[site_key]
    response = await page.goto(url, wait_until="domcontentloaded", timeout=REQUEST_TIMEOUT_MS)
    if response and response.status >= 400:
        return None
    await page.wait_for_timeout(900)

    title = await first_text(page, config["title_selectors"], timeout=6000)
    if not title:
        body_text = await first_text(page, ["body"], timeout=3000)
        h1_match = re.search(r"\n\s*([^\n]{5,140})\s*\n", body_text)
        title = h1_match.group(1).strip() if h1_match else ""
    if not title or len(title) < 3:
        return None

    description = await first_text(page, config["description_selectors"], timeout=4000)
    if not description:
        description = (await first_text(page, ["main", "body"], timeout=3000))[:400]

    price, price_source = await extract_price(page, config)

    image_url = await first_attr(page, config["image_selectors"], "src", timeout=4000)
    image_url = urljoin(config["base"], image_url) if image_url else ""

    normalized = normalize_text(f"{title} {description}")
    return {
        "site_key": site_key,
        "site": config["base"],
        "category_url": url,
        "url": url,
        "title": title.strip(),
        "description": description[:500].strip(),
        "price": price,
        "price_source": price_source,
        "image_url": image_url,
        "normalized_text": normalized,
        "tokens": tokenize(normalized),
    }


async def scrape_category(context: BrowserContext, site_key: str, cat_path: str, sem_product: asyncio.Semaphore) -> List[Dict[str, Any]]:
    cached = load_category_cache(site_key, cat_path)
    if cached is not None:
        return cached

    page = await context.new_page()
    drop_reasons = Counter()
    products: List[Dict[str, Any]] = []
    logger.info(f"📂 Скрапим категорию: {site_key}{cat_path}")

    try:
        product_urls: List[str] = []
        for page_num in range(1, CATEGORY_PAGE_LIMIT + 1):
            url = urljoin(SITE_CONFIGS[site_key]["base"], cat_path)
            if page_num > 1:
                url = f"{url}?PAGEN_1={page_num}"
            logger.info(f"   📄 Страница {page_num}: {url}")
            try:
                response = await page.goto(url, wait_until="domcontentloaded", timeout=REQUEST_TIMEOUT_MS)
                logger.info(f"   ↪ HTTP: {response.status if response else 'n/a'}")
            except Exception as exc:
                drop_reasons["category_open_error"] += 1
                logger.warning(f"⚠️ Не удалось открыть категорию {url}: {str(exc)[:120]}")
                continue
            await page.wait_for_timeout(1800)
            collected = await collect_product_urls(page, site_key, cat_path)
            if collected == [] and site_key == "attento":
                drop_reasons["no_links_detected"] += 1
            logger.info(f"   🔗 Найдено ссылок-кандидатов: {len(collected)}")
            product_urls.extend(collected)

        product_urls = list(dict.fromkeys(product_urls))
        logger.info(f"   🧾 Уникальных карточек в категории: {len(product_urls)}")

        async def scrape_one(product_url: str) -> Optional[Dict[str, Any]]:
            async with sem_product:
                product_page = await context.new_page()
                try:
                    product = await scrape_product_page(product_page, site_key, product_url)
                    if not product:
                        drop_reasons["missing_title_or_bad_card"] += 1
                        return None
                    if product["price"] <= 0:
                        drop_reasons["missing_price"] += 1
                        logger.warning(f"⚠️ Цена не распознана: {product_url} | source={product.get('price_source')}")
                        return None
                    return product
                except Exception as exc:
                    drop_reasons["product_error"] += 1
                    logger.warning(f"⚠️ Ошибка карточки {product_url}: {str(exc)[:120]}")
                    return None
                finally:
                    await product_page.close()

        results = await asyncio.gather(*(scrape_one(url) for url in product_urls), return_exceptions=True)
        for result in results:
            if isinstance(result, dict):
                products.append(result)
    finally:
        await page.close()

    unique_products = list({item["url"]: item for item in products}.values())
    save_category_cache(site_key, cat_path, unique_products)
    logger.info(f"📊 Причины отбраковки {site_key}{cat_path}: {dict(drop_reasons)}")
    logger.success(f"🎉 Категория {site_key}{cat_path} завершена: {len(unique_products)} товаров")
    return unique_products


async def scrape_categories_by_plan(
    context: BrowserContext,
    plan: Dict[str, List[str]],
    path_to_categories: Dict[Tuple[str, str], List[str]],
    sem_product: asyncio.Semaphore,
) -> Dict[str, List[Dict[str, Any]]]:
    sem_category = asyncio.Semaphore(CATEGORY_CONCURRENCY)
    category_buckets: Dict[str, List[Dict[str, Any]]] = defaultdict(list)

    async def run_category(site_key: str, cat_path: str):
        async with sem_category:
            products = await scrape_category(context, site_key, cat_path, sem_product)
            return site_key, cat_path, products

    tasks = [run_category(site_key, cat_path) for site_key, paths in plan.items() for cat_path in paths]
    results = await asyncio.gather(*tasks, return_exceptions=True)

    for result in results:
        if isinstance(result, Exception):
            logger.warning(f"⚠️ Ошибка при сборе категории: {result}")
            continue
        site_key, cat_path, products = result
        for category_key in path_to_categories.get((site_key, cat_path), []):
            category_buckets[category_key].extend(products)

    for category_key in sorted(set(CATEGORY_RULES) | set(category_buckets)):
        category_buckets[category_key] = list({item["url"]: item for item in category_buckets.get(category_key, [])}.values())
        logger.success(f"🗂️ Пул категории {category_key}: {len(category_buckets[category_key])} товаров")
    return dict(category_buckets)


# ================= РАНЖИРОВАНИЕ =================
def cheap_text_score(query: str, analog: Dict[str, Any]) -> float:
    normalized_query = normalize_text(query)
    title = analog.get("title", "")
    description = analog.get("description", "")
    haystack = analog.get("normalized_text", normalize_text(f"{title} {description}"))
    score_a = fuzz.token_set_ratio(normalized_query, haystack)
    score_b = fuzz.partial_ratio(normalized_query, title)
    query_tokens = set(tokenize(normalized_query))
    analog_tokens = set(analog.get("tokens", tokenize(haystack)))
    overlap = len(query_tokens & analog_tokens) / max(1, len(query_tokens))
    return round((score_a * 0.55 + score_b * 0.25) / 100 + overlap * 0.20, 4)


def rank_analogs(item: Dict[str, Any], analogs: List[Dict[str, Any]], analog_limit: int) -> List[Dict[str, Any]]:
    query = item["query"]
    original_photo = item.get("original_photo", "")
    logger.info(f"   🧮 Ранжируем {len(analogs)} кандидатов в категории {item['predicted_category']}...")

    for analog in analogs:
        analog["text_score"] = cheap_text_score(query, analog)

    candidates = sorted(analogs, key=lambda row: row["text_score"], reverse=True)[:TEXT_CANDIDATES_PER_ITEM]
    logger.info(f"   📊 После текстового отбора осталось: {len(candidates)}")

    for analog in candidates:
        analog["image_score"] = round(image_similarity_score(original_photo, analog.get("image_url", "")), 3)
        analog["final_score"] = round(
            analog["text_score"] * TEXT_SCORE_WEIGHT + analog["image_score"] * IMAGE_SCORE_WEIGHT,
            3,
        )

    relevant_pool = sorted(candidates, key=lambda row: row["final_score"], reverse=True)[:RELEVANCE_POOL_FOR_CHEAPEST]
    with_price = [row for row in relevant_pool if row.get("price", 0) > 0]
    selected = sorted(with_price, key=lambda row: (row["price"], -row["final_score"]))[:analog_limit]
    logger.info(f"   💰 Выбрано самых дешёвых релевантных аналогов: {len(selected)}")
    return selected


# ================= EXCEL / ИЗОБРАЖЕНИЯ =================
def ensure_text_column(df: pd.DataFrame, column_name: str) -> pd.Series:
    if column_name not in df.columns:
        logger.warning(f"⚠️ В Excel нет колонки '{column_name}' — заполняем пустыми значениями")
        return pd.Series([""] * len(df), index=df.index, dtype="object")
    return df[column_name].fillna("").astype(str)


def log_excel_diagnostics(excel_path: str, raw_df: pd.DataFrame) -> None:
    path = Path(excel_path)
    resolved_path = path.resolve() if path.is_absolute() else (Path.cwd() / path).resolve()
    exists = path.exists()
    size_mb = (path.stat().st_size / (1024 * 1024)) if exists else 0.0
    logger.info(f"🧪 Диагностика Excel: absolute={resolved_path}")
    logger.info(f"🧪 Диагностика Excel: exists={exists} | suffix={path.suffix} | size_mb={size_mb:.2f}")
    logger.info(f"🧪 Диагностика Excel: rows_raw={len(raw_df)} | cols_raw={len(raw_df.columns)}")
    logger.info(f"🧪 Диагностика Excel: columns={list(raw_df.columns)}")
    logger.info(f"🧪 Диагностика Excel: первые строки={raw_df.head(8).to_dict(orient='records')}")


def prepare_items(df: pd.DataFrame) -> pd.DataFrame:
    df = df.iloc[2:].reset_index(drop=True)
    df = df[df["Наименование"].notna() & (df["Наименование"].astype(str).str.strip() != "")].reset_index(drop=True)
    df["Описание"] = ensure_text_column(df, "Описание")
    df["Размер ш/г/в"] = ensure_text_column(df, "Размер ш/г/в")
    df["Фото_оригинала"] = ensure_text_column(df, "Фото_оригинала")
    if "Условное изображение" in df.columns:
        df["Условное изображение"] = ensure_text_column(df, "Условное изображение")
        df.loc[df["Фото_оригинала"].str.strip() == "", "Фото_оригинала"] = df.loc[df["Фото_оригинала"].str.strip() == "", "Условное изображение"]
    df["predicted_category"] = df.apply(
        lambda row: classify_item_category(str(row.get("Наименование", "")), str(row.get("Описание", ""))),
        axis=1,
    )
    df["query"] = df.apply(
        lambda row: " ".join(part for part in [str(row.get("Наименование", "")).strip(), str(row.get("Описание", "")).strip(), str(row.get("Размер ш/г/в", "")).strip()] if part),
        axis=1,
    )
    return df


def resize_image_for_excel(image_path: Path) -> Optional[Path]:
    try:
        with Image.open(image_path) as img:
            img.thumbnail((EMBED_IMAGE_WIDTH, EMBED_IMAGE_HEIGHT))
            output = IMAGE_CACHE_DIR / f"thumb_{image_path.name}.png"
            img.save(output, format="PNG")
            return output
    except Exception:
        return None


def embed_images_into_excel(output_excel: str, result_rows: List[Dict[str, Any]], analog_limit: int) -> None:
    workbook = load_workbook(output_excel)
    sheet = workbook.active
    header_map = {cell.value: idx + 1 for idx, cell in enumerate(sheet[1])}
    image_columns = ["Фото_оригинала"] + [f"Аналог_{i}_image" for i in range(1, analog_limit + 1)]
    for col_name in image_columns:
        if col_name in header_map:
            sheet.column_dimensions[sheet.cell(row=1, column=header_map[col_name]).column_letter].width = 20

    for row_idx, row in enumerate(result_rows, start=2):
        sheet.row_dimensions[row_idx].height = 95
        image_map = {"Фото_оригинала": row.get("Фото_оригинала", "")}
        for i in range(1, analog_limit + 1):
            image_map[f"Аналог_{i}_image"] = row.get(f"Аналог_{i}_image", "")
        for field_name, image_ref in image_map.items():
            if not image_ref or field_name not in header_map:
                continue
            image_path = download_to_cache(image_ref) if str(image_ref).startswith("http") else Path(str(image_ref))
            if not image_path or not image_path.exists():
                continue
            thumb = resize_image_for_excel(image_path)
            if not thumb:
                continue
            xl_img = XLImage(str(thumb))
            xl_img.width = EMBED_IMAGE_WIDTH
            xl_img.height = EMBED_IMAGE_HEIGHT
            sheet.add_image(xl_img, sheet.cell(row=row_idx, column=header_map[field_name]).coordinate)

    workbook.save(output_excel)


# ================= ОСНОВНОЙ ПОТОК =================
async def main():
    ensure_cache_dirs()
    analog_limit = ask_analog_count()
    logger.info("🚀 === СТАРТ ПОДБОРА АНАЛОГОВ ПО КАТЕГОРИЯМ ===")
    logger.info(f"📂 Читаем Excel: {EXCEL_PATH}")
    raw_df = pd.read_excel(EXCEL_PATH, sheet_name="Table 1", header=0)
    log_excel_diagnostics(EXCEL_PATH, raw_df)
    df = prepare_items(raw_df)
    logger.success(f"✅ Загружено {len(df)} товаров")

    category_stats = df["predicted_category"].value_counts().to_dict()
    logger.info(f"🧠 Авто-классификация товаров по категориям: {category_stats}")
    scrape_plan, path_to_categories = build_scrape_plan(df)
    logger.info(f"🗺️ План скрапинга: {scrape_plan}")
    logger.info(f"🧭 Карта категорий к путям: {path_to_categories}")

    category_workers, product_workers = get_optimal_workers()
    global CATEGORY_CONCURRENCY, PRODUCT_CONCURRENCY
    CATEGORY_CONCURRENCY = category_workers
    PRODUCT_CONCURRENCY = product_workers

    async with async_playwright() as pw:
        context = await create_context(pw)
        sem_product = asyncio.Semaphore(PRODUCT_CONCURRENCY)
        analog_pools = await scrape_categories_by_plan(context, scrape_plan, path_to_categories, sem_product)
        await context.close()

    all_dump_rows = []
    for category_key, items in analog_pools.items():
        for item in items:
            dump = dict(item)
            dump["predicted_category"] = category_key
            all_dump_rows.append(dump)
    if all_dump_rows:
        pd.DataFrame(all_dump_rows).to_excel(DEBUG_ANALOGS_DUMP, index=False)
        logger.info(f"🧪 Отладочный дамп аналогов сохранён: {DEBUG_ANALOGS_DUMP}")

    result_rows = []
    for idx, row in df.iterrows():
        item = {
            "name": str(row.get("Наименование", "")).strip(),
            "description": str(row.get("Описание", "")).strip(),
            "size": str(row.get("Размер ш/г/в", "")).strip(),
            "original_photo": str(row.get("Фото_оригинала", "")).strip(),
            "predicted_category": row["predicted_category"],
            "query": row["query"],
        }
        pool = [dict(entry) for entry in analog_pools.get(item["predicted_category"], [])]
        logger.info(f"\n🔄 [{idx + 1}/{len(df)}] {item['name']} → категория {item['predicted_category']} ({CATEGORY_RULES[item['predicted_category']]['label']}); пул {len(pool)}")
        best = rank_analogs(item, pool, analog_limit)

        result = {
            "№": row.get("№"),
            "Наименование": item["name"],
            "Описание оригинал": item["description"],
            "Размер": item["size"],
            "Кол-во": row.get("Кол-во"),
            "Фото_оригинала": item["original_photo"],
            "Категория": item["predicted_category"],
            "Категория_человекочитаемо": CATEGORY_RULES[item["predicted_category"]]["label"],
            "Всего_кандидатов_в_категории": len(pool),
        }
        for i, analog in enumerate(best, 1):
            result[f"Аналог_{i}_site"] = analog["site"]
            result[f"Аналог_{i}_title"] = analog["title"]
            result[f"Аналог_{i}_price"] = analog["price"]
            result[f"Аналог_{i}_url"] = analog["url"]
            result[f"Аналог_{i}_desc"] = analog["description"]
            result[f"Аналог_{i}_image_url"] = analog.get("image_url", "")
            result[f"Аналог_{i}_image"] = analog.get("image_url", "")
            result[f"Аналог_{i}_text_score"] = analog["text_score"]
            result[f"Аналог_{i}_image_match"] = analog["image_score"]
            result[f"Аналог_{i}_final_score"] = analog["final_score"]
            result[f"Аналог_{i}_price_source"] = analog.get("price_source", "")
        logger.success(f"   ✅ Товар обработан — найдено {len(best)} аналогов")
        result_rows.append(result)

    pd.DataFrame(result_rows).to_excel(OUTPUT_EXCEL, index=False)
    embed_images_into_excel(OUTPUT_EXCEL, result_rows, analog_limit)
    logger.success(f"🎉 Всё завершено! Файл сохранён: {OUTPUT_EXCEL}")
    print(f"\n🚀 Подбор завершён. Открывай файл: {OUTPUT_EXCEL}")


if __name__ == "__main__":
    asyncio.run(main())
